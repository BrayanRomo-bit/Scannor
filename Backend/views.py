from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib import colors
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from io import BytesIO
import csv
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import logout
from django.http import StreamingHttpResponse, HttpResponse, JsonResponse, FileResponse, Http404
from django.conf import settings
from django.views.decorators.csrf import csrf_exempt
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth.hashers import make_password
from django.contrib.auth import authenticate, login, get_user_model
from django.contrib import messages
from django.db import IntegrityError, models
from django.db.models import Q
from .models import RegistroInspeccion, Usuario, Empresa, ValidacionHumana, Dispositivo, IntentoConexionEdge, ModeloIA, VersionModelo
from urllib.parse import urlparse
from datetime import datetime
import requests
import json 
import secrets
import shutil
import time
import logging
import os
import zipfile
import glob

logger = logging.getLogger(__name__)

def es_url_de_dispositivo_valida(url, dispositivo):
    try:
        parsed = urlparse(url)
        return (
            parsed.scheme in ('http', 'https')
            and parsed.path.startswith('/capturas/')
        )
    except Exception:
        return False

@csrf_exempt
def api_provisionar_dispositivo(request):
    if request.method == 'POST':
        try:
            datos = json.loads(request.body)
            mac = datos.get('mac_address')
            ip_cliente = datos.get('ip_address') or request.META.get('HTTP_X_FORWARDED_FOR', '').split(',')[0].strip() or request.META.get('REMOTE_ADDR')

            if not mac:
                return JsonResponse({"error": "MAC Address no proporcionada"}, status=400)

            mac_norm = mac.strip().upper()

            try:
                dev_id = datos.get('device_id')
                if dev_id:
                    dispositivo = Dispositivo.objects.get(device_id=dev_id, activo=True)
                else:
                    dispositivos = Dispositivo.objects.filter(mac_address__iexact=mac_norm, activo=True).order_by('-id')
                    if not dispositivos.exists():
                        raise Dispositivo.DoesNotExist
                    dispositivo = dispositivos.first()
                
                IntentoConexionEdge.objects.filter(mac_address__iexact=mac_norm).delete()
                
                if ip_cliente and not dispositivo.ip_local:
                    dispositivo.ip_local = ip_cliente
                    dispositivo.save(update_fields=['ip_local'])

                host_actual = request.get_host()
                api_url = os.getenv("DJANGO_API_URL") or f"{request.scheme}://{host_actual}/api/guardar-falla/"
                
                return JsonResponse({
                    "status": "success",
                    "env_data": {
                        "SCANNOR_DEVICE_ID": dispositivo.device_id,
                        "SCANNOR_DEVICE_KEY": dispositivo.api_key,
                        "SCANNOR_API_KEY": dispositivo.admin_key,
                        "SCANNOR_STREAM_TOKEN": dispositivo.stream_token,
                        "DJANGO_API_URL": api_url
                    }
                }, status=200)

            except Dispositivo.DoesNotExist:
                IntentoConexionEdge.objects.update_or_create(
                    mac_address=mac_norm,
                    defaults={'ip_origen': ip_cliente}
                )
                
                logger.warning(f"Intento fallido. MAC registrada en Descubrimiento: {mac_norm} (IP: {ip_cliente})")
                return JsonResponse({
                    "error": "Dispositivo pendiente de autorización. Revise el panel Staff.",
                    "mac_address": mac_norm,
                    "ip_origen": ip_cliente
                }, status=404)

        except Exception as e:
            logger.error(f"Error en aprovisionamiento MAC: {e}", exc_info=True)
            return JsonResponse({"error": "Error interno del servidor"}, status=500)

    return JsonResponse({"error": "Método no permitido"}, status=405)

def es_superadmin(user):
    return user.is_authenticated and (user.is_superuser or user.rol == 'superadmin')

def es_admin_empresa(user):
    return user.is_authenticated and user.rol == 'admin_empresa'

@login_required
def descargar_env_dispositivo(request, dispositivo_id):
    dispositivo = get_object_or_404(Dispositivo, id=dispositivo_id)
    
    if not request.user.is_superuser and request.user.empresa != dispositivo.empresa:
        return HttpResponse("No tienes permisos para acceder a las credenciales de esta empresa.", status=403)
        
    host_actual = request.get_host()
    api_url = os.getenv("DJANGO_API_URL") or f"{request.scheme}://{host_actual}/api/guardar-falla/"
    
    contenido_env = f"""# ==========================================
# CONFIGURACIÓN AUTO-GENERADA SCANNORQA
# Cliente: {dispositivo.empresa.nombre}
# Dispositivo: {dispositivo.nombre}
# Fecha de generación: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
# ==========================================

SCANNOR_DEVICE_ID={dispositivo.device_id}
SCANNOR_DEVICE_KEY={dispositivo.api_key}

# --- CONFIGURACIONES GLOBALES DEL SERVIDOR ---
SCANNOR_API_KEY={dispositivo.admin_key}
SCANNOR_STREAM_TOKEN={dispositivo.stream_token}
DJANGO_API_URL={api_url}
"""
    
    response = HttpResponse(contenido_env, content_type='text/plain')
    response['Content-Disposition'] = f'attachment; filename="scannor_{dispositivo.device_id}.env"'
    
    return response

@login_required
@user_passes_test(es_superadmin)
def staff_modelos(request):
    modelos = ModeloIA.objects.all().order_by('-fecha_creacion')
    return render(request, 'staff/modelos.html', {'modelos': modelos})

@login_required
@user_passes_test(es_superadmin)
def crear_modelo_base(request):
    if request.method == 'POST':
        nombre = request.POST.get('nombre')
        tipo_industria = request.POST.get('tipo_industria')
        
        if nombre:
            ModeloIA.objects.create(nombre=nombre, tipo_industria=tipo_industria)
            messages.success(request, f"Contenedor '{nombre}' creado. Ahora puedes subirle versiones .pt.")
        else:
            messages.error(request, "El nombre del proyecto es obligatorio.")
    return redirect('staff_modelos')

@login_required
@user_passes_test(es_superadmin)
def subir_version_modelo(request):
    if request.method == 'POST':
        modelo_id = request.POST.get('modelo_id')
        version = request.POST.get('version')
        archivo = request.FILES.get('archivo')
        
        if modelo_id and version and archivo:
            if not archivo.name.endswith('.pt'):
                messages.error(request, "Error: El archivo debe ser un modelo de PyTorch (.pt).")
                return redirect('staff_modelos')
            
            modelo = get_object_or_404(ModeloIA, id=modelo_id)
            try:
                VersionModelo.objects.create(
                    modelo=modelo,
                    version=version,
                    archivo=archivo,
                    activa=True 
                )
                messages.success(request, f"Versión {version} desplegada en '{modelo.nombre}'.")
            except IntegrityError:
                messages.error(request, f"Error: La versión {version} ya existe para este modelo.")
        else:
            messages.error(request, "Faltan datos para subir la versión.")
    return redirect('staff_modelos')

@login_required
@user_passes_test(es_superadmin)
def activar_version(request, version_id):
    version = get_object_or_404(VersionModelo, id=version_id)
    version.activa = True
    version.save()
    messages.success(request, f"Rollback exitoso. La versión {version.version} vuelve a estar en producción.")
    return redirect('staff_modelos')

@login_required
@user_passes_test(es_superadmin)
def eliminar_modelo(request, modelo_id):
    modelo = get_object_or_404(ModeloIA, id=modelo_id)
    modelo.delete()
    messages.success(request, "Proyecto de IA y todo su historial de versiones eliminados.")
    return redirect('staff_modelos')

@csrf_exempt
def api_descargar_modelo(request):
    api_key = request.headers.get("X-Device-Auth")
    dispositivo = Dispositivo.objects.filter(api_key=api_key).first()

    if not dispositivo:
        return JsonResponse({"error": "No autorizado. API Key inválida."}, status=403)
        
    if not dispositivo.modelo_asignado:
        return JsonResponse({"error": "Esta cámara no tiene un proyecto de IA asignado."}, status=404)

    version_act = dispositivo.modelo_asignado.version_activa
    
    if not version_act or not version_act.archivo:
        return JsonResponse({"error": "El proyecto asignado no tiene ninguna versión .pt activa."}, status=404)

    ruta_archivo = version_act.archivo.path
    
    if os.path.exists(ruta_archivo):
        response = FileResponse(open(ruta_archivo, 'rb'), content_type='application/octet-stream')
        response['Content-Disposition'] = f'attachment; filename="modelo_{dispositivo.modelo_asignado.id}_{version_act.version}.pt"'
        response['X-Model-Version'] = str(version_act.version)
        return response
    else:
        return JsonResponse({"error": "El archivo físico del modelo desapareció del servidor."}, status=500)

@user_passes_test(es_superadmin)
def staff_empresas(request):
    if request.method == 'POST':
        nombre_empresa = request.POST.get('nombre')
        limite = request.POST.get('limite_usuarios', 5)
        
        admin_username = request.POST.get('admin_username')
        admin_email = request.POST.get('admin_email', '')
        admin_password = request.POST.get('admin_password')
        
        if nombre_empresa and admin_username and admin_password:
            if Empresa.objects.filter(nombre=nombre_empresa).exists():
                messages.error(request, f"La empresa '{nombre_empresa}' ya está registrada.")
                return redirect('staff_empresas')
                
            if Usuario.objects.filter(username=admin_username).exists():
                messages.error(request, f"El usuario '@{admin_username}' ya está en uso.")
                return redirect('staff_empresas')
                
            nueva_empresa = Empresa.objects.create(
                nombre=nombre_empresa,
                limite_usuarios=int(limite),
                estado_plan='activo'
            )
            
            Usuario.objects.create_user(
                username=admin_username,
                email=admin_email,
                password=admin_password,
                empresa=nueva_empresa,
                rol='admin_empresa',
                is_active=True
            )
            
            messages.success(request, f"¡Fábrica '{nombre_empresa}' y Admin '@{admin_username}' creados exitosamente!")
            return redirect('staff_empresas')

    empresas_activas = Empresa.objects.all().order_by('-id')
    return render(request, 'staff/empresas.html', {'empresas': empresas_activas})

@login_required
@user_passes_test(es_superadmin)
def staff_dispositivos(request):
    if request.method == 'POST':
        nombre = request.POST.get('nombre')
        device_id = request.POST.get('device_id')
        mac_address = request.POST.get('mac_address')
        ip_local = request.POST.get('ip_local', '').strip()
        empresa_id = request.POST.get('empresa_id')
        modelo_id = request.POST.get('modelo_id')

        if nombre and device_id and mac_address and empresa_id:
            try:
                empresa = Empresa.objects.get(id=empresa_id)
                modelo_obj = ModeloIA.objects.get(id=modelo_id) if modelo_id else None
                
                Dispositivo.objects.create(
                    nombre=nombre,
                    device_id=device_id,
                    mac_address=mac_address.upper(), 
                    ip_local=ip_local,
                    empresa=empresa,
                    modelo_asignado=modelo_obj
                )
                
                IntentoConexionEdge.objects.filter(mac_address__iexact=mac_address).delete()
                
                messages.success(request, f"Dispositivo '{nombre}' registrado exitosamente.")
            except Empresa.DoesNotExist:
                messages.error(request, "La empresa seleccionada no existe.")
            except Exception as e:
                messages.error(request, f"Error al registrar dispositivo: Verifica que el ID Técnico '{device_id}' no esté duplicado.")
        else:
            messages.error(request, "Todos los campos principales son obligatorios.")
            
        return redirect('staff_dispositivos')

    dispositivos = Dispositivo.objects.all().order_by('-id')
    empresas = Empresa.objects.all()
    intentos_conexion = IntentoConexionEdge.objects.all().order_by('-ultimo_intento')
    modelos_ia = ModeloIA.objects.all().order_by('-id') 

    context = {
        'dispositivos': dispositivos,
        'empresas': empresas,
        'intentos_conexion': intentos_conexion,
        'modelos': modelos_ia
    }
    return render(request, 'staff/dispositivos.html', context)

@login_required
@user_passes_test(es_superadmin)
def editar_dispositivo(request, dispositivo_id):
    dispositivo = get_object_or_404(Dispositivo, id=dispositivo_id)
    
    if request.method == 'POST':
        dispositivo.nombre = request.POST.get('nombre')
        dispositivo.device_id = request.POST.get('device_id')
        dispositivo.mac_address = request.POST.get('mac_address')
        dispositivo.ip_local = request.POST.get('ip_local')
        
        empresa_id = request.POST.get('empresa_id')
        if empresa_id: dispositivo.empresa_id = empresa_id
            
        modelo_id = request.POST.get('modelo_id')
        if modelo_id:
            dispositivo.modelo_asignado_id = modelo_id
            dispositivo.modelo_operativo = True
        else:
            dispositivo.modelo_asignado = None
            dispositivo.modelo_operativo = False
            
        dispositivo.save()
        messages.success(request, f"Dispositivo '{dispositivo.nombre}' actualizado correctamente.")
        return redirect('staff_dispositivos')
        
    empresas = Empresa.objects.all()
    modelos = ModeloIA.objects.all()
    return render(request, 'staff/editar_dispositivo.html', {
        'dispositivo': dispositivo, 'empresas': empresas, 'modelos': modelos
    })

@login_required
@user_passes_test(es_superadmin)
def extraer_dataset(request, dispositivo_id):
    dispositivo = get_object_or_404(Dispositivo, id=dispositivo_id)

    if not dispositivo.ip_local:
        messages.error(request, "El dispositivo no tiene una IP configurada.")
        return redirect('staff_dispositivos')

    ip_edge = dispositivo.ip_con_puerto if hasattr(dispositivo, 'ip_con_puerto') else dispositivo.ip_local
    url_edge = f"http://{ip_edge}/api/descargar-dataset" if "://" not in ip_edge else f"{ip_edge}/api/descargar-dataset"

    try:
        respuesta = requests.get(url_edge, headers={"X-API-Key": dispositivo.admin_key}, stream=True, timeout=15)

        if respuesta.status_code == 200:
            response = HttpResponse(respuesta.content, content_type='application/zip')
            response['Content-Disposition'] = f'attachment; filename="Dataset_{dispositivo.nombre}.zip"'
            return response
        else:
            messages.error(request, f"La cámara rechazó la conexión. Código: {respuesta.status_code}")
    except requests.exceptions.RequestException:
        messages.error(request, f"No se pudo conectar a la cámara en {dispositivo.ip_local}. ¿Está encendida?")

    return redirect('staff_dispositivos')

@login_required
@user_passes_test(es_superadmin)
def eliminar_dispositivo(request, dispositivo_id):
    dispositivo = get_object_or_404(Dispositivo, id=dispositivo_id)
    nombre = dispositivo.nombre
    dispositivo.delete()
    messages.success(request, f"Cámara '{nombre}' eliminada permanentemente.")
    return redirect('staff_dispositivos')

Usuario = get_user_model()
@user_passes_test(es_superadmin)
def staff_usuarios(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        email = request.POST.get('email')
        password = request.POST.get('password')
        empresa_id = request.POST.get('empresa_id')

        if username and password and empresa_id:
            try:
                empresa_obj = Empresa.objects.get(id=empresa_id)
                Usuario.objects.create_user(
                    username=username,
                    email=email,
                    password=password,
                    rol='admin_empresa',
                    empresa=empresa_obj,
                    is_active=True
                )
                messages.success(request, f"Gerente {username} creado exitosamente para la empresa {empresa_obj.nombre}.")
            except IntegrityError:
                messages.error(request, f"Error: El usuario '{username}' ya existe en el sistema.")
            except Empresa.DoesNotExist:
                messages.error(request, "Error: La empresa seleccionada no existe.")
                
        return redirect('staff_usuarios')

    usuarios = Usuario.objects.filter(is_superuser=False).select_related('empresa').order_by('-date_joined')
    empresas = Empresa.objects.filter(activa=True)
    
    return render(request, 'staff/usuarios.html', {
        'usuarios': usuarios,
        'empresas': empresas
    })

@user_passes_test(es_superadmin)
def toggle_empresa(request, empresa_id):
    empresa = get_object_or_404(Empresa, id=empresa_id)
    empresa.activa = not empresa.activa
    
    if empresa.activa:
        empresa.estado_plan = 'activo'
    else:
        empresa.estado_plan = 'suspendido'
        
    empresa.save()
    return redirect('staff_empresas')

@user_passes_test(es_superadmin)
def toggle_usuario(request, usuario_id):
    usuario = get_object_or_404(Usuario, id=usuario_id)
    if not usuario.is_superuser: 
        usuario.is_active = not usuario.is_active
        usuario.save()
    return redirect('staff_usuarios')

@login_required
def gestion_planta(request):
    if request.user.is_superuser or request.user.rol == 'superadmin':
        return redirect('staff_empresas')
    if not es_admin_empresa(request.user):
        return redirect('dashboard')
    empresa_actual = request.user.empresa
    operadores = Usuario.objects.filter(empresa=empresa_actual, rol='operador')
    usuarios_activos = Usuario.objects.filter(empresa=empresa_actual, is_active=True).count()
    
    dispositivos_empresa = Dispositivo.objects.filter(empresa=empresa_actual, activo=True)
    
    return render(request, 'paneles/gestion_planta.html', {
        'empresa': empresa_actual,
        'operadores': operadores,
        'usuarios_activos': usuarios_activos,
        'limite': empresa_actual.limite_usuarios,
        'dispositivos': dispositivos_empresa
    })

@login_required
def crear_operador(request):
    if request.user.is_superuser or request.user.rol == 'superadmin':
        return redirect('staff_empresas')
    if not es_admin_empresa(request.user):
        return redirect('dashboard')
    if request.method == 'POST':
        empresa_actual = request.user.empresa
        
        usuarios_activos = Usuario.objects.filter(empresa=empresa_actual, is_active=True).count()
        if usuarios_activos >= empresa_actual.limite_usuarios:
            messages.error(request, f"¡Límite alcanzado! Tu plan permite un máximo de {empresa_actual.limite_usuarios} usuarios.")
            return redirect('gestion_planta')
            
        nuevo_username = request.POST.get('username')
        nuevo_email = request.POST.get('email')
        nueva_password = request.POST.get('password')
        
        if nuevo_username and nueva_password:
            if Usuario.objects.filter(username=nuevo_username).exists():
                messages.error(request, f"El usuario @{nuevo_username} ya está en uso. Elige otro.")
                return redirect('gestion_planta')
                
            Usuario.objects.create_user(
                username=nuevo_username,
                email=nuevo_email,
                password=nueva_password, 
                empresa=empresa_actual,
                rol='operador',
                is_active=True 
            )
            messages.success(request, f"Operador @{nuevo_username} creado exitosamente.")
            
    return redirect('gestion_planta')

def login_usuario(request):
    if request.user.is_authenticated:
        if request.user.is_superuser or request.user.rol == 'superadmin':
            return redirect('staff_empresas')
        else:
            return redirect('dashboard')

    if request.method == 'POST':
        usuario_ingresado = (request.POST.get('username') or '').strip()
        password_ingresada = request.POST.get('password') or ''
        
        usuario_obj = Usuario.objects.filter(username__iexact=usuario_ingresado).first()
        if usuario_obj:
            usuario_ingresado = usuario_obj.username

        user = authenticate(request, username=usuario_ingresado, password=password_ingresada)
        
        if user is not None:
            if user.is_active:
                login(request, user)
                
                if user.is_superuser or user.rol == 'superadmin':
                    return redirect('staff_empresas')
                else:
                    return redirect('dashboard')
            else:
                return redirect('sala_espera')
        else:
            error = "Usuario o contraseña incorrectos."
            return render(request, 'paneles/login.html', {'error': error})
            
    return render(request, 'paneles/login.html')

def sala_espera(request):
    return render(request, 'paneles/sala_espera.html')

def cerrar_sesion(request):
    logout(request)
    return redirect('login')
    
@login_required
def dashboard(request):
    if request.user.is_superuser or request.user.rol == 'superadmin':
        return redirect('staff_empresas')

    if hasattr(request.user, 'empresa') and request.user.empresa:
        dispositivos = Dispositivo.objects.filter(empresa=request.user.empresa, activo=True)
    else:
        dispositivos = Dispositivo.objects.none()

    dispositivo_activo = dispositivos.first()
    es_modelo_ajedrez = dispositivo_activo.es_modelo_ajedrez if dispositivo_activo else False
    modelos_disponibles = ModeloIA.objects.all().order_by('nombre')

    context = {
        'dispositivos': dispositivos,
        'dispositivo_activo': dispositivo_activo,
        'es_modelo_ajedrez': es_modelo_ajedrez,
        'modelos_disponibles': modelos_disponibles,
        'admin_key': dispositivo_activo.admin_key if dispositivo_activo else '',
        'stream_token': dispositivo_activo.stream_token if dispositivo_activo else '',
    }
    
    return render(request, 'base.html', context)

@login_required
@csrf_exempt
def api_dispositivo_control(request):
    """
    Relay/Proxy transparente que permite a operadores y administradores de la empresa
    controlar la máquina Edge (Máquina 1: iniciar, detener, consultar estado, escanear hardware),
    incluso si la interfaz se consulta desde otra máquina, tablet o red diferente.
    """
    if hasattr(request.user, 'empresa') and request.user.empresa:
        dispositivo = Dispositivo.objects.filter(empresa=request.user.empresa, activo=True).first()
    elif request.user.is_superuser or request.user.rol == 'superadmin':
        dispositivo = Dispositivo.objects.filter(activo=True).first()
    else:
        return JsonResponse({'error': 'No tienes permisos para operar estaciones'}, status=403)

    if not dispositivo:
        return JsonResponse({'error': 'No hay estación de trabajo activa vinculada a tu empresa'}, status=404)

    default_port = int(os.getenv("PORT_DEFECT_INSPECTION", "8001"))
    default_host = os.getenv("DEFAULT_EDGE_HOST", "127.0.0.1")
    puerto = dispositivo.puerto_servicio if hasattr(dispositivo, 'puerto_servicio') else default_port
    candidatos_url = []

    base_cfg = dispositivo.ip_con_puerto if hasattr(dispositivo, 'ip_con_puerto') else (dispositivo.ip_local or f"{default_host}:{puerto}")
    if "://" not in base_cfg:
        base_cfg = f"http://{base_cfg}"
    candidatos_url.append(base_cfg.rstrip('/'))

    for fallback in [f"http://{default_host}:{puerto}", f"http://localhost:{puerto}"]:
        if fallback not in candidatos_url:
            candidatos_url.append(fallback)

    try:
        from microservices.shared.edge_credentials import obtener_ip_local
        lan_ip = obtener_ip_local()
        if lan_ip:
            lan_url = f"http://{lan_ip}:{puerto}"
            if lan_url not in candidatos_url:
                candidatos_url.append(lan_url)
    except Exception:
        pass

    if request.method == 'POST':
        try:
            data = json.loads(request.body.decode('utf-8')) if request.body else {}
        except Exception:
            data = request.POST.dict()
        action = data.get('action') or request.POST.get('action')
        payload = data.get('payload')
    else:
        action = request.GET.get('action', 'health')
        payload = None

    if not action:
        return JsonResponse({'error': 'Parámetro action requerido'}, status=400)

    method_map = {
        'health': 'GET',
        'start': 'POST',
        'stop': 'POST',
        'clear': 'DELETE',
        'estado-camaras': 'GET',
        'escanear-hardware': 'GET',
        'guardar-configuracion': 'POST',
    }

    http_method = method_map.get(action, 'POST' if request.method == 'POST' else 'GET')
    headers = {
        'X-API-Key': dispositivo.admin_key,
        'Content-Type': 'application/json'
    }

    resp = None
    ultimo_error = None
    url_exitosa = None

    for base_url in candidatos_url:
        target_url = f"{base_url}/api/{action}"
        to = 2 if len(candidatos_url) > 1 and base_url != candidatos_url[-1] else 5
        try:
            if http_method == 'GET':
                resp = requests.get(target_url, headers=headers, timeout=to)
                if resp.status_code == 404:
                    resp = requests.get(f"{base_url}/{action}", headers=headers, timeout=to)
            elif http_method == 'POST':
                resp = requests.post(target_url, headers=headers, json=payload if payload else {}, timeout=to)
                if resp.status_code == 404:
                    resp = requests.post(f"{base_url}/{action}", headers=headers, json=payload if payload else {}, timeout=to)
            elif http_method == 'DELETE':
                resp = requests.delete(target_url, headers=headers, timeout=to)
                if resp.status_code == 404:
                    resp = requests.delete(f"{base_url}/{action}", headers=headers, timeout=to)
            else:
                return JsonResponse({'error': 'Método no soportado'}, status=405)

            if resp is not None and resp.status_code < 500:
                url_exitosa = base_url
                break
        except requests.exceptions.RequestException as e:
            ultimo_error = e
            continue

    if resp is None:
        return JsonResponse({
            'error': f"No se pudo contactar a la estación Edge en {candidatos_url[0]}. Verifica que la Máquina 1 esté encendida con el motor activo en el puerto {puerto}.",
            'detail': str(ultimo_error)
        }, status=502)

    if url_exitosa and url_exitosa != candidatos_url[0]:
        try:
            nuevo_host = url_exitosa.split("://")[-1]
            dispositivo.ip_local = nuevo_host
            dispositivo.save(update_fields=['ip_local'])
        except Exception:
            pass

    try:
        data_resp = resp.json()
        if action == 'health' and isinstance(data_resp, dict):
            if 'motor_activo' not in data_resp and data_resp.get('status') == 'healthy':
                data_resp['motor_activo'] = True
        return JsonResponse(data_resp, status=resp.status_code, safe=False)
    except Exception:
        return HttpResponse(resp.text, status=resp.status_code, content_type=resp.headers.get('content-type', 'application/json'))

@login_required
def api_dispositivo_video_stream(request, cam_id):
    """
    Proxy de streaming de video MJPEG para permitir visualización en tiempo real
    cuando el operador está en una máquina cliente que no puede acceder directamente al puerto de la cámara.
    """
    if hasattr(request.user, 'empresa') and request.user.empresa:
        dispositivo = Dispositivo.objects.filter(empresa=request.user.empresa, activo=True).first()
    elif request.user.is_superuser or request.user.rol == 'superadmin':
        dispositivo = Dispositivo.objects.filter(activo=True).first()
    else:
        return HttpResponse("No autorizado", status=403)

    if not dispositivo:
        return HttpResponse("Estación no configurada", status=404)

    default_port = int(os.getenv("PORT_DEFECT_INSPECTION", "8001"))
    default_host = os.getenv("DEFAULT_EDGE_HOST", "127.0.0.1")
    puerto = dispositivo.puerto_servicio if hasattr(dispositivo, 'puerto_servicio') else default_port
    candidatos_url = []
    base_cfg = dispositivo.ip_con_puerto if hasattr(dispositivo, 'ip_con_puerto') else (dispositivo.ip_local or f"{default_host}:{puerto}")
    if "://" not in base_cfg:
        base_cfg = f"http://{base_cfg}"
    candidatos_url.append(base_cfg.rstrip('/'))

    for fallback in [f"http://{default_host}:{puerto}", f"http://localhost:{puerto}"]:
        if fallback not in candidatos_url:
            candidatos_url.append(fallback)

    r = None
    ultimo_error = None
    for base_url in candidatos_url:
        url = f"{base_url}/video_feed/{cam_id}?token={dispositivo.stream_token}"
        try:
            r = requests.get(url, stream=True, timeout=2)
            if r.status_code == 200:
                break
        except Exception as e:
            ultimo_error = e
            continue

    if r is not None and r.status_code == 200:
        return StreamingHttpResponse(
            r.iter_content(chunk_size=4096),
            content_type=r.headers.get('content-type', 'multipart/x-mixed-replace; boundary=frame')
        )
    return HttpResponse(f"Error conectando a stream de cámara ({cam_id}): {str(ultimo_error)}", status=502)

@login_required
@user_passes_test(es_superadmin)
def cambiar_modelo_estacion(request):
    if request.method == 'POST':
        dispositivo_id = request.POST.get('dispositivo_id')
        modelo_id = request.POST.get('modelo_id')
        dispositivo = get_object_or_404(Dispositivo, id=dispositivo_id)
            
        if modelo_id:
            modelo = get_object_or_404(ModeloIA, id=modelo_id)
            dispositivo.modelo_asignado = modelo
            dispositivo.modelo_operativo = True
            messages.success(request, f"Estación '{dispositivo.nombre}' vinculada a: {modelo.nombre}")
        else:
            dispositivo.modelo_asignado = None
            dispositivo.modelo_operativo = False
            messages.info(request, f"Estación '{dispositivo.nombre}' configurada sin modelo IA.")
            
        dispositivo.save()
    return redirect(request.META.get('HTTP_REFERER', 'dashboard'))

@login_required
def personalizacion(request):
    if request.user.is_superuser or request.user.rol == 'superadmin':
        return redirect('staff_empresas')

    dispositivos = Dispositivo.objects.none()
    api_key_local = ""
    default_host = os.getenv("DEFAULT_EDGE_HOST", "127.0.0.1")
    ip_edge = default_host

    if hasattr(request.user, 'empresa') and request.user.empresa:
        dispositivos = Dispositivo.objects.filter(empresa=request.user.empresa, activo=True)

    primero = None
    if dispositivos.exists():
        primero = dispositivos.first()
        api_key_local = primero.admin_key
        ip_edge = primero.ip_local or default_host

    modelos_disponibles = ModeloIA.objects.all().order_by('nombre')
    context = {
        'dispositivos': dispositivos,
        'dispositivo_activo': primero,
        'api_key': api_key_local,
        'ip_edge': ip_edge,
        'modelos_disponibles': modelos_disponibles,
    }
    return render(request, 'paneles/personalizacion.html', context)

@csrf_exempt 
def recibir_falla_api(request):
    if request.method == 'POST':
        try:
            datos = json.loads(request.body)
            campos_requeridos = ['uid', 'dispositivo_id', 'timestamp', 'lote', 'falla', 'confianza', 'estado', 'imagen']
            if not all(k in datos for k in campos_requeridos):
                return JsonResponse({"error": "Payload incompleto"}, status=400)
            
            try:
                dispositivo_obj = Dispositivo.objects.get(device_id=datos['dispositivo_id'])
            except Dispositivo.DoesNotExist:
                ip_cliente = request.META.get('REMOTE_ADDR')
                IntentoConexionEdge.objects.update_or_create(
                    mac_address=datos['dispositivo_id'], 
                    defaults={'ip_origen': ip_cliente}
                )
                logger.warning(f"Cámara no autorizada detectada en el Radar: {datos['dispositivo_id']}")
                return JsonResponse({"error": f"Dispositivo técnico '{datos['dispositivo_id']}' no encontrado. Revisa el radar Staff."}, status=404)
            
            auth_header = request.headers.get('X-Device-Auth', '')
            if not secrets.compare_digest(auth_header, dispositivo_obj.api_key):
                logger.warning(f"Intento de acceso denegado para cámara: {dispositivo_obj.nombre}")
                return JsonResponse({"error": "No autorizado: API Key incorrecta"}, status=403)

            if not es_url_de_dispositivo_valida(datos['imagen'], dispositivo_obj):
                return JsonResponse({"error": "URL de imagen rechazada (Políticas SSRF)."}, status=400)

            dir_capturas = os.path.join(settings.MEDIA_ROOT, 'capturas_inspecciones')
            os.makedirs(dir_capturas, exist_ok=True)
            fname_base = os.path.basename(urlparse(datos['imagen']).path) or f"{datos['uid']}.jpg"
            fname_limpio = fname_base.replace('_vis.jpg', '.jpg')
            fname_vis = fname_limpio.replace('.jpg', '_vis.jpg')

            img_b64 = datos.get('imagen_base64')
            if img_b64:
                try:
                    if ',' in img_b64:
                        img_b64 = img_b64.split(',')[1]
                    img_bytes = base64.b64decode(img_b64)
                    with open(os.path.join(dir_capturas, fname_vis), 'wb') as f_img:
                        f_img.write(img_bytes)
                except Exception as ex_b64:
                    logger.warning(f"No se pudo guardar imagen visual localmente en Django: {ex_b64}")

            img_limpia_b64 = datos.get('imagen_limpia_base64')
            if img_limpia_b64:
                try:
                    if ',' in img_limpia_b64:
                        img_limpia_b64 = img_limpia_b64.split(',')[1]
                    img_clean_bytes = base64.b64decode(img_limpia_b64)
                    with open(os.path.join(dir_capturas, fname_limpio), 'wb') as f_clean:
                        f_clean.write(img_clean_bytes)
                except Exception as ex_clean:
                    logger.warning(f"No se pudo guardar imagen limpia localmente en Django: {ex_clean}")
            elif img_b64 and not fname_base.endswith('_vis.jpg'):
                try:
                    with open(os.path.join(dir_capturas, fname_limpio), 'wb') as f_clean:
                        f_clean.write(img_bytes)
                except Exception:
                    pass

            if datos.get('bbox_txt'):
                txt_name = fname_limpio.replace('.jpg', '.txt')
                with open(os.path.join(dir_capturas, txt_name), 'w', encoding='utf-8') as f_txt:
                    f_txt.write(datos['bbox_txt'])

            fecha_utc = datetime.fromisoformat(datos['timestamp'])
            
            inspeccion, creado = RegistroInspeccion.objects.update_or_create(
                uid=datos['uid'],
                defaults={
                    'dispositivo': dispositivo_obj,
                    'empresa': dispositivo_obj.empresa,
                    'timestamp': fecha_utc,
                    'lote_qr': datos['lote'],
                    'falla': datos['falla'],
                    'confianza': float(datos['confianza']),
                    'estado': datos['estado'],
                    'imagen': datos['imagen']
                }
            )
            
            dispositivo_obj.ultimo_sync = fecha_utc
            dispositivo_obj.save(update_fields=['ultimo_sync'])
            
            return JsonResponse({"status": "Guardado exitosamente", "uid": inspeccion.uid}, status=201)
            
        except ValueError as ve:
            return JsonResponse({"error": f"Error de formato: {str(ve)}"}, status=400)
        except Exception as e:
            logger.error(f"Error interno en DB: {e}", exc_info=True)
            return JsonResponse({"error": "Error interno del servidor"}, status=500)
            
    return JsonResponse({"error": "Método no permitido"}, status=405)

def api_servir_imagen_inspeccion(request, uid):
    """
    Sirve la evidencia fotográfica de una inspección con tolerancia a equipos offline.
    1. Busca la imagen en disco local (MEDIA_ROOT, capturas_falla, capturas_test).
    2. Si es URL remota, intenta obtenerla por red y cachearla en Django.
    3. Si el equipo Edge está offline o la imagen física no existe, genera una tarjeta
       vectorial SVG con la telemetría exacta (Lote, Estado, Falla, Confianza) evitando imágenes rotas.
    """
    inspeccion = get_object_or_404(RegistroInspeccion, uid=uid)
    nombre_archivo = os.path.basename(urlparse(inspeccion.imagen).path) if inspeccion.imagen else f"{uid}.jpg"

    posibles_rutas = [
        os.path.join(settings.MEDIA_ROOT, 'capturas_inspecciones', nombre_archivo),
        os.path.join(settings.MEDIA_ROOT, 'capturas_inspecciones', nombre_archivo.replace('.jpg', '_vis.jpg')),
        os.path.join(settings.BASE_DIR, 'capturas_falla', nombre_archivo),
        os.path.join(settings.BASE_DIR, 'capturas_falla', nombre_archivo.replace('.jpg', '_vis.jpg')),
        os.path.join(settings.BASE_DIR, 'scratch', 'capturas_test', nombre_archivo),
        os.path.join(settings.BASE_DIR, 'scratch', 'capturas_test', nombre_archivo.replace('.jpg', '_vis.jpg')),
    ]

    for ruta in posibles_rutas:
        if os.path.isfile(ruta):
            return FileResponse(open(ruta, 'rb'), content_type='image/jpeg')

    curados = glob.glob(os.path.join(settings.MEDIA_ROOT, 'dataset_curado', '**', nombre_archivo), recursive=True)
    if curados and os.path.isfile(curados[0]):
        return FileResponse(open(curados[0], 'rb'), content_type='image/jpeg')

    if inspeccion.imagen and inspeccion.imagen.startswith(('http://', 'https://')):
        try:
            res = requests.get(inspeccion.imagen, timeout=2.0)
            if res.status_code == 200:
                dir_cache = os.path.join(settings.MEDIA_ROOT, 'capturas_inspecciones')
                os.makedirs(dir_cache, exist_ok=True)
                with open(os.path.join(dir_cache, nombre_archivo), 'wb') as f_cache:
                    f_cache.write(res.content)
                return HttpResponse(res.content, content_type='image/jpeg')
        except Exception:
            pass

    es_ok = 'aprob' in inspeccion.estado.lower()
    color_bg = "#091e17" if es_ok else "#240a0e"
    color_stroke = "#10b981" if es_ok else "#ef4444"
    tag_bg = "#059669" if es_ok else "#dc2626"
    tag_label = "APROBADO" if es_ok else "RECHAZADO"

    svg_content = f'''<svg xmlns="http://www.w3.org/2000/svg" width="640" height="360" viewBox="0 0 640 360">
        <defs>
            <linearGradient id="grad" x1="0%" y1="0%" x2="100%" y2="100%">
                <stop offset="0%" style="stop-color:{color_bg};stop-opacity:1" />
                <stop offset="100%" style="stop-color:#0b1329;stop-opacity:1" />
            </linearGradient>
        </defs>
        <rect width="640" height="360" fill="url(#grad)" rx="12"/>
        <rect x="12" y="12" width="616" height="336" fill="none" stroke="{color_stroke}" stroke-width="2" stroke-dasharray="8 6" rx="8"/>
        <text x="320" y="95" font-family="-apple-system, BlinkMacSystemFont, Segoe UI, Roboto, sans-serif" font-size="20" font-weight="700" fill="#f8fafc" text-anchor="middle">
            Evidencia de Inspección Edge (Offline)
        </text>
        <rect x="230" y="125" width="180" height="34" fill="{tag_bg}" rx="17"/>
        <text x="320" y="148" font-family="-apple-system, BlinkMacSystemFont, Segoe UI, Roboto, sans-serif" font-size="13" font-weight="800" fill="#ffffff" text-anchor="middle">
            {tag_label}: {inspeccion.falla}
        </text>
        <text x="320" y="200" font-family="-apple-system, BlinkMacSystemFont, Segoe UI, Roboto, sans-serif" font-size="13" font-weight="600" fill="#94a3b8" text-anchor="middle">
            Lote: {inspeccion.lote_qr}  ·  Confianza IA: {inspeccion.confianza*100:.1f}%
        </text>
        <text x="320" y="235" font-family="-apple-system, BlinkMacSystemFont, Segoe UI, Roboto, sans-serif" font-size="12" fill="#64748b" text-anchor="middle">
            Dispositivo: {inspeccion.dispositivo.nombre if inspeccion.dispositivo else "Edge"}  ·  {inspeccion.timestamp.strftime('%Y-%m-%d %H:%M:%S')}
        </text>
        <text x="320" y="280" font-family="-apple-system, BlinkMacSystemFont, Segoe UI, Roboto, sans-serif" font-size="11" fill="#475569" text-anchor="middle">
            Registro sincronizado desde buffer local SQLite tras recuperación de red.
        </text>
    </svg>'''
    return HttpResponse(svg_content, content_type='image/svg+xml')

@login_required
def api_obtener_resultados(request):
    if request.user.is_superuser or request.user.rol == 'superadmin':
        inspecciones = RegistroInspeccion.objects.select_related('dispositivo', 'empresa').filter(archivado=False).order_by('-timestamp')
    elif request.user.empresa:
        inspecciones = RegistroInspeccion.objects.select_related('dispositivo', 'empresa').filter(empresa=request.user.empresa, archivado=False).order_by('-timestamp')
    else:
        inspecciones = RegistroInspeccion.objects.none()
    data = []
    for insp in inspecciones:
        data.append({
            'uid': insp.uid,
            'dispositivo': insp.dispositivo.nombre if insp.dispositivo else 'N/A',
            'timestamp': insp.timestamp.strftime('%Y-%m-%d %H:%M:%S'),
            'lote': insp.lote_qr,
            'falla': insp.falla,
            'confianza': f"{insp.confianza * 100:.1f}%" if insp.confianza <= 1 else f"{insp.confianza:.1f}%",
            'estado': insp.estado,
            'imagen': f"/api/inspeccion/imagen/{insp.uid}/"
        })
    
    return JsonResponse({'resultados': data})

@csrf_exempt
@login_required
def registrar_falso_positivo(request, uid):
    if request.method == 'POST':
        try:
            inspeccion = RegistroInspeccion.objects.select_related('dispositivo__modelo_asignado').get(uid=uid)
        except RegistroInspeccion.DoesNotExist:
            return JsonResponse({"error": "La pieza aún se está sincronizando. Espera 5 segundos."}, status=404)
        
        if not request.user.is_superuser and inspeccion.empresa != request.user.empresa:
            return JsonResponse({"error": "No tienes permiso para alterar esta pieza"}, status=403)
            
        try:
            data = json.loads(request.body)
            justificacion_texto = data.get('justificacion', 'Sin justificación')
            es_acierto = data.get('modelo_acerto', False)
            tipo_auditoria = data.get('tipo_auditoria', 'rechazado')
            nuevo_defecto = data.get('nuevo_defecto', '')
            
            modelo = inspeccion.dispositivo.modelo_asignado if inspeccion.dispositivo else None
            industria = modelo.tipo_industria if modelo else 'general'

            if tipo_auditoria == 'aprobado':
                if es_acierto:
                    inspeccion.estado = 'Aprobado (Auditado)'
                    inspeccion.falla = 'bloque_ok'
                    falla_real = 'bloque_ok'
                    categoria_dataset = 'aprobados'
                    falla_limpia = 'bloque_ok'
                else:
                    defecto_final = nuevo_defecto if nuevo_defecto else 'bloque_fisura'
                    inspeccion.estado = 'Rechazado (Manual)'
                    inspeccion.falla = defecto_final
                    falla_real = defecto_final
                    categoria_dataset = 'rechazados'
                    falla_limpia = defecto_final.replace(" ", "_").lower()
            else:
                if es_acierto:
                    inspeccion.estado = 'Rechazado (Auditado)'
                    falla_real = inspeccion.falla
                    categoria_dataset = 'rechazados'
                    falla_limpia = inspeccion.falla.replace(" ", "_").lower()
                else:
                    inspeccion.estado = 'Aprobado (Manual)'
                    falla_real = 'falso_positivo'
                    categoria_dataset = 'falsos_positivos'
                    falla_limpia = 'falsos_positivos'

            validacion, created = ValidacionHumana.objects.update_or_create(
                inspeccion=inspeccion,
                defaults={
                    'validado_por': request.user,
                    'modelo_acerto': es_acierto,
                    'falla_real': falla_real,
                    'justificacion': justificacion_texto
                }
            )

            ruta_destino = os.path.join(settings.MEDIA_ROOT, 'dataset_curado', industria, categoria_dataset, falla_limpia)
            os.makedirs(ruta_destino, exist_ok=True)
            
            nombre_base = os.path.basename(urlparse(inspeccion.imagen).path)
            if not nombre_base:
                nombre_base = f"{uid}.jpg"
                
            nombre_limpio = nombre_base.replace('_vis.jpg', '.jpg')
            url_img_limpia = inspeccion.imagen.replace('_vis.jpg', '.jpg')
            url_txt = url_img_limpia.replace('.jpg', '.txt')

            ruta_img = os.path.join(ruta_destino, nombre_limpio)
            ruta_txt = os.path.join(ruta_destino, nombre_limpio.replace('.jpg', '.txt'))
            
            posibles_locales = [
                os.path.join(settings.MEDIA_ROOT, 'capturas_inspecciones', nombre_limpio),
                os.path.join(settings.BASE_DIR, 'capturas_falla', nombre_limpio),
                os.path.join(settings.BASE_DIR, 'scratch', 'capturas_test', nombre_limpio),
            ]
            if not nombre_base.endswith('_vis.jpg'):
                posibles_locales.append(os.path.join(settings.MEDIA_ROOT, 'capturas_inspecciones', nombre_base))
                posibles_locales.append(os.path.join(settings.BASE_DIR, 'capturas_falla', nombre_base))
            
            imagen_extraida = False
            for p_local in posibles_locales:
                if os.path.isfile(p_local):
                    try:
                        shutil.copy2(p_local, ruta_img)
                        txt_local = p_local.replace('.jpg', '.txt')
                        if os.path.isfile(txt_local):
                            shutil.copy2(txt_local, ruta_txt)
                        else:
                            with open(ruta_txt, 'w', encoding='utf-8') as f:
                                f.write(f"# Auditoria Scannor QA: {falla_real}\n# Lote: {inspeccion.lote_qr}\n")
                        imagen_extraida = True
                        break
                    except Exception as ex_copy:
                        logger.warning(f"Error copiando evidencia local para dataset: {ex_copy}")

            if not imagen_extraida and es_url_de_dispositivo_valida(url_img_limpia, inspeccion.dispositivo):
                try:
                    res_img = requests.get(url_img_limpia, timeout=3)
                    if res_img.status_code != 200:
                        res_img = requests.get(inspeccion.imagen, timeout=3)

                    if res_img.status_code == 200:
                        with open(ruta_img, 'wb') as f:
                            f.write(res_img.content)
                            
                        res_txt = requests.get(url_txt, timeout=2)
                        if res_txt.status_code == 200:
                            with open(ruta_txt, 'wb') as f:
                                f.write(res_txt.content)
                        else:
                            with open(ruta_txt, 'w', encoding='utf-8') as f:
                                f.write(f"# Auditoria Scannor QA: {falla_real}\n# Lote: {inspeccion.lote_qr}\n")
                                
                        imagen_extraida = True
                except Exception:
                    pass

            validacion.curado_para_dataset = imagen_extraida
            validacion.save()
            inspeccion.save()

            if imagen_extraida:
                return JsonResponse({"status": "Auditoría guardada y dataset extraído exitosamente.", "curado_para_dataset": True})
            else:
                return JsonResponse({"status": "Auditoría registrada exitosamente en el historial. (La cámara estuvo offline; la validación humana quedó firmada en base de datos).", "curado_para_dataset": False})
            
        except Exception as e:
            return JsonResponse({"error": f"Error interno: {str(e)}"}, status=500)
            
    return JsonResponse({"error": "Método inválido"}, status=405)

@csrf_exempt
@login_required
def api_archivar_sesion(request):
    if request.method == 'POST':
        if request.user.is_superuser:
            RegistroInspeccion.objects.filter(archivado=False).update(archivado=True)
        elif request.user.empresa:
            RegistroInspeccion.objects.filter(empresa=request.user.empresa, archivado=False).update(archivado=True)
        return JsonResponse({"status": "Sesión archivada."})
    return JsonResponse({"error": "Método no permitido"}, status=405)

@csrf_exempt
@login_required
def api_borrar_todas_imagenes(request):
    if request.method == 'POST':
        total_borrados = 0
        
        ruta_capturas = os.path.join(settings.BASE_DIR, 'capturas_falla')
        if os.path.exists(ruta_capturas):
            for archivo in os.listdir(ruta_capturas):
                ruta_f = os.path.join(ruta_capturas, archivo)
                try:
                    if os.path.isfile(ruta_f) or os.path.islink(ruta_f):
                        os.unlink(ruta_f)
                        total_borrados += 1
                except Exception as e:
                    logger.error(f"Error borrando archivo de evidencia {ruta_f}: {e}")

        ruta_dataset = os.path.join(settings.MEDIA_ROOT, 'dataset_curado')
        if os.path.exists(ruta_dataset):
            for root, dirs, files in os.walk(ruta_dataset, topdown=False):
                for f in files:
                    fp = os.path.join(root, f)
                    try:
                        os.unlink(fp)
                        total_borrados += 1
                    except Exception as e:
                        logger.error(f"Error borrando dataset curado {fp}: {e}")
                for d in dirs:
                    dp = os.path.join(root, d)
                    try:
                        if not os.listdir(dp):
                            os.rmdir(dp)
                    except Exception:
                        pass

        ValidacionHumana.objects.filter(curado_para_dataset=True).update(curado_para_dataset=False)

        if request.user.is_superuser:
            RegistroInspeccion.objects.filter(archivado=False).update(archivado=True)
        elif hasattr(request.user, 'empresa') and request.user.empresa:
            RegistroInspeccion.objects.filter(empresa=request.user.empresa, archivado=False).update(archivado=True)

        offline_db = os.path.join(settings.BASE_DIR, 'scannor_offline.db')
        if os.path.exists(offline_db):
            try:
                import sqlite3
                conn = sqlite3.connect(offline_db)
                conn.execute("DELETE FROM inspecciones WHERE sync = 1")
                conn.commit()
                conn.close()
            except Exception as e:
                logger.warning(f"Error limpiando scannor_offline.db: {e}")

        try:
            if request.user.is_superuser:
                dispositivos = Dispositivo.objects.filter(activo=True)
            elif hasattr(request.user, 'empresa') and request.user.empresa:
                dispositivos = Dispositivo.objects.filter(empresa=request.user.empresa, activo=True)
            else:
                dispositivos = []

            for disp in dispositivos:
                if disp.ip_local and disp.ip_local not in ('127.0.0.1', 'localhost'):
                    host = disp.ip_con_puerto
                    try:
                        requests.delete(
                            f"http://{host}/api/clear",
                            headers={"X-API-Key": disp.admin_key},
                            timeout=1.5
                        )
                    except Exception:
                        pass
        except Exception as e:
            logger.warning(f"Error notificando a dispositivos edge para borrado de capturas: {e}")

        return JsonResponse({
            "status": "success",
            "mensaje": f"Se eliminaron {total_borrados} archivos de imagen/evidencia y dataset curado del disco.",
            "total_borrados": total_borrados
        })
    return JsonResponse({"error": "Método no permitido"}, status=405)

@login_required
def descargar_instalador(request, os_type):
    if os_type == 'windows':
        filename = "ScannorQA_Edge_Setup.exe"
        content = "ESTE ES UN ARCHIVO PLACEHOLDER SIMULANDO EL EJECUTABLE.\nEn el siguiente paso usaremos PyInstaller para compilar el motor y ponerlo aquí."
        response = HttpResponse(content, content_type='application/octet-stream')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
    else:
        filename = "ScannorQA_Motor_Linux.tar.gz"
        file_path = os.path.join(settings.BASE_DIR, 'static', filename)
        
        if os.path.exists(file_path):
            response = FileResponse(open(file_path, 'rb'), as_attachment=True, filename=filename)
            return response
        else:
            raise Http404("El instalador de Linux no se encontró en el servidor.")

@login_required
@user_passes_test(es_superadmin)
def push_update_ota(request, dispositivo_id):
    dispositivo = get_object_or_404(Dispositivo, id=dispositivo_id)
    
    if not dispositivo.ip_local:
        messages.error(request, "El dispositivo no tiene IP configurada.")
        return redirect('staff_dispositivos')

    ip_edge = dispositivo.ip_con_puerto if hasattr(dispositivo, 'ip_con_puerto') else dispositivo.ip_local
    url_edge = f"http://{ip_edge}/api/actualizar-modelo-ahora" if "://" not in ip_edge else f"{ip_edge}/api/actualizar-modelo-ahora"
    
    try:
        respuesta = requests.post(url_edge, headers={"X-API-Key": dispositivo.admin_key}, timeout=30)
        
        if respuesta.status_code == 200:
            messages.success(request, f"¡Actualización OTA en caliente exitosa para la cámara '{dispositivo.nombre}'!")
        else:
            detalles = respuesta.json().get('detail', 'Error desconocido')
            messages.error(request, f"La cámara falló la recarga OTA. Detalle: {detalles}")
    except requests.exceptions.RequestException:
        messages.error(request, f"No se pudo contactar a {dispositivo.ip_local}. ¿El motor Edge está corriendo en esa máquina?")
        
    return redirect('staff_dispositivos')

@login_required
@user_passes_test(es_superadmin)
def descargar_dataset_curado(request):
    filtro = request.GET.get('filtro', 'ambos').lower().strip()
    ruta_dataset = os.path.join(settings.MEDIA_ROOT, 'dataset_curado')

    qs_validaciones = ValidacionHumana.objects.filter(
        curado_para_dataset=True,
        inspeccion__archivado=False
    ).select_related('inspeccion', 'inspeccion__dispositivo', 'inspeccion__dispositivo__modelo_asignado')

    if not request.user.is_superuser and hasattr(request.user, 'empresa') and request.user.empresa:
        qs_validaciones = qs_validaciones.filter(inspeccion__empresa=request.user.empresa)

    if not qs_validaciones.exists() or not os.path.exists(ruta_dataset):
        messages.error(request, "El dataset curado de la sesión activa está vacío. Audita piezas en el panel antes de descargar.")
        return redirect('staff_modelos')

    try:
        buffer = BytesIO()
        archivos_incluidos = 0
        arcnames_agregados = set()

        with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            for val in qs_validaciones:
                if val.falla_real == 'falso_positivo' or val.inspeccion.estado == 'Aprobado (Manual)':
                    continue

                es_aprobado = val.falla_real == 'bloque_ok' or val.inspeccion.estado.startswith('Aprobado')
                es_rechazado = not es_aprobado

                if filtro == 'aprobados' and not es_aprobado:
                    continue
                if filtro == 'rechazados' and not es_rechazado:
                    continue

                base_fname = os.path.basename(urlparse(val.inspeccion.imagen).path).replace('_vis.jpg', '.jpg')
                txt_fname = base_fname.replace('.jpg', '.txt')
                cat_dir = 'aprobados' if es_aprobado else 'rechazados'
                falla_dir = val.falla_real.replace(' ', '_').lower() if val.falla_real else 'general'

                for root, dirs, files in os.walk(ruta_dataset):
                    if base_fname in files:
                        full_p = os.path.join(root, base_fname)
                        arcname = os.path.join(cat_dir, falla_dir, base_fname)
                        if arcname not in arcnames_agregados:
                            zip_file.write(full_p, arcname=arcname)
                            arcnames_agregados.add(arcname)
                            archivos_incluidos += 1
                    if txt_fname in files:
                        full_p = os.path.join(root, txt_fname)
                        arcname = os.path.join(cat_dir, falla_dir, txt_fname)
                        if arcname not in arcnames_agregados:
                            zip_file.write(full_p, arcname=arcname)
                            arcnames_agregados.add(arcname)
                            archivos_incluidos += 1

        if archivos_incluidos == 0:
            etiqueta = "aprobadas (conformes)" if filtro == 'aprobados' else ("rechazadas (con defectos)" if filtro == 'rechazados' else "")
            messages.error(request, f"No se encontraron muestras {etiqueta} válidas en la sesión activa (falsos positivos excluidos).")
            return redirect('staff_modelos')

        buffer.seek(0)
        nombres_filtro = {'aprobados': 'Aprobados', 'rechazados': 'Rechazados', 'ambos': 'Completo'}
        nom_sufijo = nombres_filtro.get(filtro, 'Completo')
        filename = f"Dataset_Curado_{nom_sufijo}_{int(time.time())}.zip"

        response = HttpResponse(buffer.getvalue(), content_type='application/zip')
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
    except Exception as e:
        messages.error(request, f"Error al generar el ZIP: {e}")
        return redirect('staff_modelos')

@csrf_exempt
def api_edge_status(request):
    if request.method == 'POST':
        try:
            datos = json.loads(request.body)
            device_id = datos.get('dispositivo_id')
            modelo_instalado = datos.get('modelo_instalado', False)

            api_key = request.headers.get("X-Device-Auth")
            dispositivo = Dispositivo.objects.filter(device_id=device_id, api_key=api_key).first()

            if not dispositivo:
                return JsonResponse({"error": "No autorizado o dispositivo no encontrado."}, status=403)
            
            dispositivo.modelo_operativo = modelo_instalado
            dispositivo.save(update_fields=['modelo_operativo'])
                
            return JsonResponse({"status": "Estado actualizado exitosamente"}, status=200)
        except Exception as e:
            logger.error(f"Error actualizando estado Edge: {e}")
            return JsonResponse({"error": "Error interno"}, status=500)
    return JsonResponse({"error": "Método no permitido"}, status=405)

def _resultados_para_exportar(request):
    if request.user.is_superuser or request.user.rol == 'superadmin':
        qs = RegistroInspeccion.objects.select_related('dispositivo', 'empresa').filter(archivado=False)
        titulo_empresa = "Todas las plantas"
    elif request.user.empresa:
        qs = RegistroInspeccion.objects.select_related('dispositivo', 'empresa').filter(
            empresa=request.user.empresa, archivado=False
        )
        titulo_empresa = request.user.empresa.nombre
    else:
        return RegistroInspeccion.objects.none(), "N/A"

    dispositivo_id = request.GET.get('dispositivo_id')
    if dispositivo_id:
        qs = qs.filter(dispositivo_id=dispositivo_id)

    estado = request.GET.get('estado')
    if estado:
        qs = qs.filter(estado=estado)

    return qs.order_by('-timestamp'), titulo_empresa

@login_required
def exportar_pdf_resultados(request):
    inspecciones, titulo_empresa = _resultados_para_exportar(request)
    inspecciones = inspecciones[:2000]

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(letter),
        leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.5 * cm, bottomMargin=1.5 * cm,
    )
    estilos = getSampleStyleSheet()
    elementos = [
        Paragraph(f"Reporte de Inspecciones — {titulo_empresa}", estilos['Title']),
        Paragraph(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", estilos['Normal']),
        Spacer(1, 0.5 * cm),
    ]

    filas = [['UID', 'Dispositivo', 'Fecha', 'Lote', 'Falla', 'Confianza', 'Estado']]
    for insp in inspecciones:
        confianza_pct = insp.confianza * 100 if insp.confianza <= 1 else insp.confianza
        filas.append([
            insp.uid,
            insp.dispositivo.nombre if insp.dispositivo else 'N/A',
            insp.timestamp.strftime('%Y-%m-%d %H:%M'),
            insp.lote_qr or '-',
            insp.falla,
            f"{confianza_pct:.1f}%",
            insp.estado,
        ])

    tabla = Table(filas, repeatRows=1)
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f2937')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f3f4f6')]),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elementos.append(tabla)
    doc.build(elementos)
    buffer.seek(0)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="reporte_inspecciones_{int(time.time())}.pdf"'
    return response

@login_required
def exportar_csv_resultados(request):
    inspecciones, _ = _resultados_para_exportar(request)

    response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
    response['Content-Disposition'] = f'attachment; filename="reporte_inspecciones_{int(time.time())}.csv"'

    writer = csv.writer(response)
    writer.writerow(['UID', 'Dispositivo', 'Fecha', 'Lote', 'Falla', 'Confianza (%)', 'Estado', 'Imagen'])
    for insp in inspecciones.iterator():
        confianza_pct = insp.confianza * 100 if insp.confianza <= 1 else insp.confianza
        writer.writerow([
            insp.uid,
            insp.dispositivo.nombre if insp.dispositivo else 'N/A',
            insp.timestamp.strftime('%Y-%m-%d %H:%M:%S'),
            insp.lote_qr or '',
            insp.falla,
            f"{confianza_pct:.1f}",
            insp.estado,
            insp.imagen,
        ])
    return response

@login_required
def exportar_excel_resultados(request):
    inspecciones, _ = _resultados_para_exportar(request)

    wb = Workbook()
    ws = wb.active
    ws.title = "Inspecciones"

    ws.append(['UID', 'Dispositivo', 'Fecha', 'Lote', 'Falla', 'Confianza (%)', 'Estado', 'Imagen'])
    for celda in ws[1]:
        celda.font = Font(bold=True, color='FFFFFF')
        celda.fill = PatternFill(start_color='1F2937', end_color='1F2937', fill_type='solid')
        celda.alignment = Alignment(horizontal='center')

    for insp in inspecciones.iterator():
        confianza_pct = insp.confianza * 100 if insp.confianza <= 1 else insp.confianza
        ws.append([
            insp.uid,
            insp.dispositivo.nombre if insp.dispositivo else 'N/A',
            insp.timestamp.strftime('%Y-%m-%d %H:%M:%S'),
            insp.lote_qr or '',
            insp.falla,
            round(confianza_pct, 1),
            insp.estado,
            insp.imagen,
        ])

    for col, ancho in zip('ABCDEFGH', [22, 20, 18, 14, 18, 14, 18, 40]):
        ws.column_dimensions[col].width = ancho

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="reporte_inspecciones_{int(time.time())}.xlsx"'
    return response